import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=90, bottom=90, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_bw(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11.5)
    r_body.font.italic = italic
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_laporan_p6(output_filename):
    doc = docx.Document()

    # Margin 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # HEADER / TITLE (Pure Black & White)
    # ---------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("LAPORAN PRAKTIKUM PEMROGRAMAN BERORIENTASI OBYEK")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r = p_sub.add_run("MODUL 6: POLYMORPHISM, METHOD OVERRIDING, OVERLOADING, DAN DYNAMIC BINDING")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    # Identity Table (Black & White)
    table_id = doc.add_table(rows=4, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False
    
    id_data = [
        ("Nama Mahasiswa", ": Rafael Rizky"),
        ("NRP", ": 3125522007"),
        ("Dosen Pengampu", ": Nirwana Haidar Hari, S.Pd., M.Kom."),
        ("Program Studi / Kampus", ": D3 Teknik Informatika - PENS PSDKU Sumenep (2026)")
    ]
    for idx, (label, val) in enumerate(id_data):
        row = table_id.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0, 0, 0)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------------------------
    # HALAMAN 1 / BAGIAN 1: REVIEW P5, SPRINT PLANNING & AUDIT BEHAVIOR
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "1. Deskripsi Proyek dan Review Hasil P5", level=1)
    add_body_p(doc, "Sistem Kasir Sederhana", bold_prefix="1.1 Nama Proyek: ")
    add_body_p(doc, "Pada praktikum P5, hierarki inheritance telah berhasil dibangun dengan Superclass Produk serta Subclass ProdukMakanan dan ProdukElektronik. Namun, pemanggilan method pada subclass masih bersifat statis dan belum memanfaatkan keunggulan Polymorphism. Pada operasional kasir sesungguhnya, sistem kasir perlu memproses dan menampilkan beraneka ragam produk secara seragam melalui tipe referensi induk (Produk), namun masing-masing produk harus mampu merespons dengan perilaku yang kontekstual (dynamic method dispatch), seperti membedakan pencetakan tanggal kadaluarsa untuk produk konsumsi dan jaminan garansi untuk produk elektronik. Selain itu, transaksi pembayaran kasir membutuhkan variasi pemanggilan method (overloading) untuk memfasilitasi transaksi standar maupun transaksi pembayaran tunai yang memerlukan perhitungan uang kembalian. Oleh karena itu, pada praktikum P6 ini diterapkan konsep Polymorphism, Method Overriding, Method Overloading, dan Dynamic Binding.", bold_prefix="1.2 Latar Belakang Masalah: ")

    add_styled_heading(doc, "2. Perencanaan Agile Sprint (Sprint P6)", level=1)
    add_body_p(doc, "Mengembangkan behavior object pada proyek Sistem Kasir Sederhana melalui method overriding (@Override) dan polymorphism sehingga subclass dapat merespons pemanggilan method yang sama dengan perilaku berbeda secara dinamis di runtime, serta menyediakan fleksibilitas input pembayaran melalui method overloading.", bold_prefix="2.1 Sprint Goal: ")

    add_body_p(doc, "Sprint Backlog P6", bold_prefix="2.2 ")
    t_sb = doc.add_table(rows=8, cols=3)
    t_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_sb)
    sb_data = [
        ("ID", "Sprint Backlog", "Status"),
        ("SB-01", "Audit hierarki Superclass dan Subclass hasil P5.", "DONE"),
        ("SB-02", "Menentukan behavior method yang dapat dioverride pada hierarki Produk.", "DONE"),
        ("SB-03", "Mengimplementasikan method overriding (@Override) pada ProdukMakanan dan ProdukElektronik.", "DONE"),
        ("SB-04", "Mengimplementasikan method overloading pada class Transaksi (tambahItem dan prosesTransaksi).", "DONE"),
        ("SB-05", "Mengimplementasikan upcasting dan polymorphic reference.", "DONE"),
        ("SB-06", "Membuat dan mengiterasi polymorphic collection (array of Superclass Produk[]).", "DONE"),
        ("SB-07", "Melakukan pengujian pembuktian dynamic binding saat program dijalankan.", "DONE")
    ]
    for r_idx, row_content in enumerate(sb_data):
        for c_idx, val in enumerate(row_content):
            cell = t_sb.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(0.9)
            elif c_idx == 1:
                cell.width = Inches(4.7)
            else:
                cell.width = Inches(0.9)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0 or c_idx == 2:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "3. Audit Hierarchy P5 dan Identifikasi Behavior Polymorphic (Bagian A)", level=1)
    add_body_p(doc, "Berdasarkan hierarki class hasil P5 (Superclass Produk dengan Subclass ProdukMakanan dan ProdukElektronik), diidentifikasi behavior yang dimiliki superclass namun memerlukan implementasi spesifik pada masing-masing subclass:")

    t_audit = doc.add_table(rows=3, cols=4)
    t_audit.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_audit)
    audit_data = [
        ("Superclass", "Subclass", "Method yang Dioverride", "Perilaku Spesifik pada Subclass"),
        ("Produk", "ProdukMakanan", "tampilkanData(), getKategoriInfo()", "Menampilkan kategori pangan & mencetak tanggal kadaluarsa (expired date)."),
        ("Produk", "ProdukElektronik", "tampilkanData(), getKategoriInfo()", "Menampilkan kategori hardware & mencetak masa perlindungan garansi unit.")
    ]
    for r_idx, row_content in enumerate(audit_data):
        for c_idx, val in enumerate(row_content):
            cell = t_audit.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.3)
            elif c_idx == 1:
                cell.width = Inches(1.5)
            elif c_idx == 2:
                cell.width = Inches(1.8)
            else:
                cell.width = Inches(2.2)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 2 / BAGIAN 2: UML CLASS DIAGRAM, OVERRIDING, OVERLOADING & DYNAMIC BINDING
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "4. Update Class Diagram (Bagian B)", level=1)
    add_body_p(doc, "Struktur relasi dan polymorphism digambarkan melalui UML Class Diagram monokrom (hitam-putih) yang menampilkan anotasi overriding <<Override>>, method overloading <<Overload>>, serta relasi Association, Aggregation, dan Composition:")

    diagram_path = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\assets", "class_diagram_p6.png")
    if os.path.exists(diagram_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(diagram_path, width=Inches(6.2))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run("Gambar 1. UML Class Diagram Polymorphism, Overriding, dan Overloading (P6)")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0, 0, 0)

    add_styled_heading(doc, "5. Pembahasan Konsep Polymorphism Modul 6", level=1)
    
    add_body_p(doc, "Method Overriding terjadi ketika subclass mendefinisikan ulang method yang telah dideklarasikan pada superclass dengan nama, return type, dan parameter yang identik. Pada class ProdukMakanan dan ProdukElektronik, method tampilkanData() dan getKategoriInfo() dioverride menggunakan anotasi @Override. Hal ini memastikan integritas penulisan method saat proses kompilasi.", bold_prefix="5.1 Penerapan Method Overriding (@Override): ")
    
    add_body_p(doc, "Method Overloading terjadi ketika beberapa method di dalam class yang sama memiliki nama yang persis sama namun berbeda daftar parameternya (jumlah atau tipe parameter). Pada class Transaksi diterapkan dua bentuk overloading yang sangat relevan secara bisnis kasir:\n"
                           "• tambahItem(Produk, int): Penambahan barang belanja reguler.\n"
                           "• tambahItem(Produk, int, double): Penambahan barang dengan potongan diskon promosi khusus barang.\n"
                           "• prosesTransaksi(): Checkout otomatis / non-tunai standar.\n"
                           "• prosesTransaksi(double uangDiterima): Pembayaran kasir tunai yang menghitung selisih uang kembalian (change refund) atau menampilkan peringatan uang kurang.", bold_prefix="5.2 Penerapan Method Overloading: ")

    add_body_p(doc, "Upcasting adalah proses memperlakukan objek subclass sebagai tipe superclass-nya (contoh: 'Produk p = new ProdukMakanan(...)'). Ketika method yang dioverride dipanggil melalui referensi superclass ('p.tampilkanData()'), JVM akan menentukan implementasi method mana yang dieksekusi saat program berjalan (runtime) berdasarkan tipe objek nyata di memori heap, bukan berdasarkan tipe deklarasi variabel referensinya. Mekanisme ini dikenal sebagai Dynamic Binding atau Dynamic Method Dispatch.", bold_prefix="5.3 Upcasting dan Dynamic Binding: ")

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 3 / BAGIAN 3: PENGUJIAN RUNNING, TABEL DYNAMIC BINDING, REVIEW & RETROSPECTIVE
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "6. Pengujian Program dan Pembuktian Dynamic Binding", level=1)
    add_body_p(doc, "Pengujian pada Main.java mencakup eksekusi upcasting, pemanggilan method polimorfik, iterasi polymorphic collection (array of Superclass Produk[]), serta eksekusi method overloading transaksi kasir.")

    add_body_p(doc, "Tangkapan Layar Hasil Running Terminal (Monokrom)", bold_prefix="6.1 ")
    path_term = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\assets", "hasil_running.png")
    if os.path.exists(path_term):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.paragraph_format.space_before = Pt(2)
        p_img3.paragraph_format.space_after = Pt(2)
        run_img3 = p_img3.add_run()
        run_img3.add_picture(path_term, width=Inches(5.9))

        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap3.paragraph_format.space_before = Pt(0)
        p_cap3.paragraph_format.space_after = Pt(6)
        r_cap3 = p_cap3.add_run("Gambar 2. Bukti Running Program Main.java pada Terminal Windows PowerShell")
        r_cap3.font.name = "Times New Roman"
        r_cap3.font.size = Pt(9.5)
        r_cap3.font.italic = True
        r_cap3.font.color.rgb = RGBColor(0, 0, 0)

    add_body_p(doc, "Tabel Pembuktian Pengujian Dynamic Binding (Bagian F)", bold_prefix="6.2 ")
    t_db = doc.add_table(rows=5, cols=4)
    t_db.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_db)
    db_data = [
        ("Reference Type", "Actual Object", "Method Dipanggil", "Perilaku / Output yang Dihasilkan"),
        ("Produk (Superclass)", "ProdukMakanan", "tampilkanData()", "Format Makanan (Mencetak Tanggal Kadaluarsa: 25-10-2026)"),
        ("Produk (Superclass)", "ProdukElektronik", "tampilkanData()", "Format Elektronik (Mencetak Masa Garansi: 6 Bulan)"),
        ("Produk (Superclass)", "ProdukMakanan", "getKategoriInfo()", "Mengembalikan: 'Makanan & Minuman Segar (Konsumsi)'"),
        ("Produk (Superclass)", "ProdukElektronik", "getKategoriInfo()", "Mengembalikan: 'Elektronik & Aksesoris Gadget (Hardware)'")
    ]
    for r_idx, row_content in enumerate(db_data):
        for c_idx, val in enumerate(row_content):
            cell = t_db.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.8)
            elif c_idx == 1:
                cell.width = Inches(1.5)
            elif c_idx == 2:
                cell.width = Inches(1.4)
            else:
                cell.width = Inches(2.1)
            set_cell_margins(cell, 40, 40, 60, 60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "7. Sprint Review (Evaluasi Sprint P6)", level=1)
    t_rev = doc.add_table(rows=9, cols=2)
    t_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_rev)
    rev_data = [
        ("Item Evaluasi", "Hasil / Keterangan"),
        ("Superclass dan subclass tersedia", "Superclass Produk serta Subclass ProdukMakanan dan ProdukElektronik tersedia."),
        ("Method overriding berhasil", "Method tampilkanData() dan getKategoriInfo() berhasil dioverride dengan @Override."),
        ("Minimal 2 subclass memiliki behavior berbeda", "ProdukMakanan dan ProdukElektronik memiliki format dan informasi spesifik."),
        ("Method overloading tersedia", "Method tambahItem dan prosesTransaksi berhasil dioverload dengan parameter berbeda."),
        ("Upcasting berhasil", "Referensi Produk p1..p4 berhasil merujuk objek subclass."),
        ("Polymorphic collection berhasil", "Array Produk[] berhasil menampung 4 objek subclass dan diiterasi seragam."),
        ("Dynamic binding dapat dibuktikan", "Terbukti eksekusi method ditentukan oleh objek aktual di runtime."),
        ("Program berjalan", "Program sukses dikompilasi (javac) dan dieksekusi (java) 100% tanpa error.")
    ]
    for r_idx, row_content in enumerate(rev_data):
        for c_idx, val in enumerate(row_content):
            cell = t_rev.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(2.3)
            else:
                cell.width = Inches(4.2)
            set_cell_margins(cell, 40, 40, 60, 60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "8. Sprint Retrospective", level=1)
    add_body_p(doc, "Penerapan runtime polymorphism dan dynamic binding berjalan sempurna. Struktur kode menjadi sangat fleksibel dan penambahan variasi pembayaran kasir melalui overloading memberikan nilai fungsional nyata.", bold_prefix="What Went Well? ")
    add_body_p(doc, "Membedakan konsep overloading (compile-time) dan overriding (runtime) menuntut pemahaman mendalam tentang signature method dan resolusi JVM.", bold_prefix="What Went Wrong? ")
    add_body_p(doc, "Pada sprint berikutnya (P7: Abstraction & Interface), Superclass Produk dapat ditransformasikan menjadi Abstract Class atau mengimplementasikan Interface kontrak bisnis kasir.", bold_prefix="Improvement: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "9. Definition of Done (DoD) Checklist Modul 6", level=1)
    dod_items = [
        ("Menggunakan project P5 (Sistem Kasir Sederhana)", True),
        ("Minimal 1 superclass digunakan (Produk)", True),
        ("Minimal 2 subclass digunakan (ProdukMakanan dan ProdukElektronik)", True),
        ("Terdapat method pada superclass yang dioverride (tampilkanData, getKategoriInfo)", True),
        ("Minimal 2 subclass melakukan overriding", True),
        ("Menggunakan annotation @Override secara tepat", True),
        ("Terdapat minimal 1 contoh overloading (tambahItem dan prosesTransaksi)", True),
        ("Terdapat penggunaan upcasting (Produk p = new ProdukMakanan)", True),
        ("Terdapat polymorphic reference yang dibuktikan di kode", True),
        ("Minimal 3 object digunakan dalam pengujian polymorphism (4 objek digunakan)", True),
        ("Dynamic binding berhasil ditunjukkan dan dibuktikan via tabel evaluasi", True),
        ("Class Diagram diperbarui dengan simbol polymorphic dan overloading", True),
        ("Program berhasil dikompilasi bebas error", True),
        ("Program berhasil dijalankan dengan output yang sesuai", True),
        ("Sprint Backlog diperbarui dengan status DONE", True),
        ("Sprint Review dan Sprint Retrospective tersedia lengkap", True)
    ]
    t_dod = doc.add_table(rows=len(dod_items)+1, cols=2)
    t_dod.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_dod)
    
    hdr0, hdr1 = t_dod.rows[0].cells[0], t_dod.rows[0].cells[1]
    hdr0.width, hdr1.width = Inches(5.3), Inches(1.2)
    set_cell_margins(hdr0, 50, 50, 70, 70)
    set_cell_margins(hdr1, 50, 50, 70, 70)
    p = hdr0.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Kriteria Definition of Done (Modul 6)"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 0, 0)
    p = hdr1.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Status"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 0, 0)

    for idx, (item, fulfilled) in enumerate(dod_items, start=1):
        c0, c1 = t_dod.rows[idx].cells[0], t_dod.rows[idx].cells[1]
        c0.width, c1.width = Inches(5.3), Inches(1.2)
        set_cell_margins(c0, 35, 35, 60, 60)
        set_cell_margins(c1, 35, 35, 60, 60)
        
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0); p0.paragraph_format.line_spacing = 1.1
        r0 = p0.add_run(item); r0.font.name = "Times New Roman"; r0.font.size = Pt(9); r0.font.color.rgb = RGBColor(0, 0, 0)
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0); p1.paragraph_format.line_spacing = 1.1
        r1 = p1.add_run("[ v ] Terpenuhi" if fulfilled else "[   ] Belum"); r1.font.name = "Times New Roman"; r1.font.size = Pt(9); r1.font.bold = True; r1.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # LAMPIRAN: SOURCE CODE JAVA LENGKAP P6 (Hitam & Putih)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "Lampiran: Source Code Java (Polymorphism P6)", level=1)
    
    code_files = [
        ("Produk.java (Superclass)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\Produk.java"),
        ("ProdukMakanan.java (Subclass 1, @Override)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\ProdukMakanan.java"),
        ("ProdukElektronik.java (Subclass 2, @Override)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\ProdukElektronik.java"),
        ("Pelanggan.java (Class Relasi P4)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\Pelanggan.java"),
        ("ItemTransaksi.java (Class Relasi P4/P5)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\ItemTransaksi.java"),
        ("Transaksi.java (Class Relasi P4, Overloading)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\Transaksi.java"),
        ("Main.java (Main Runner Test Polymorphism)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\src\Main.java")
    ]

    for fname, fpath in code_files:
        add_styled_heading(doc, f"Berkas: {fname}", level=2)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                code_content = f.read()
            
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(6.5)
            set_cell_margins(cell, 70, 70, 90, 90)
            set_table_borders_bw(tbl)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(code_content)
            r.font.name = "Courier New"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0, 0, 0)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_filename)
    print(f"P6 Document saved successfully: {output_filename}")

out_docx = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\Laporan_P6_3125522007_Rafael_Rizky.docx"
build_laporan_p6(out_docx)

readme_docx = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P6_3125522007_Rafael_Rizky\README.docx"
build_laporan_p6(readme_docx)

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=90, bottom=90, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_bw(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11.5)
    r_body.font.italic = italic
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_laporan_p5(output_filename):
    doc = docx.Document()

    # Margin 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # HEADER / TITLE (Pure Black & White)
    # ---------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("LAPORAN PRAKTIKUM PEMROGRAMAN BERORIENTASI OBYEK")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r = p_sub.add_run("MODUL 5: INHERITANCE, GENERALIZATION, SUPERCLASS, DAN SUBCLASS")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    # Identity Table (Black & White)
    table_id = doc.add_table(rows=4, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False
    
    id_data = [
        ("Nama Mahasiswa", ": Rafael Rizky"),
        ("NRP", ": 3125522007"),
        ("Dosen Pengampu", ": Nirwana Haidar Hari, S.Pd., M.Kom."),
        ("Program Studi / Kampus", ": D3 Teknik Informatika - PENS PSDKU Sumenep (2026)")
    ]
    for idx, (label, val) in enumerate(id_data):
        row = table_id.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0, 0, 0)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------------------------
    # HALAMAN 1 / BAGIAN 1: REVIEW P4, SPRINT PLANNING & AUDIT CLASS P4
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "1. Deskripsi Proyek dan Review Hasil P4", level=1)
    add_body_p(doc, "Sistem Kasir Sederhana", bold_prefix="1.1 Nama Proyek: ")
    add_body_p(doc, "Pada praktikum P4, relasi antarobjek (Association, Aggregation, dan Composition) telah berhasil dihubungkan menjadi satu kesatuan alur transaksi kasir. Namun, seiring perluasan variasi barang yang dijual pada toko ritel, timbul kebutuhan spesifik untuk mengelola komoditas yang berbeda karakteristiknya, seperti produk makanan (yang memerlukan pelacakan tanggal kadaluarsa) dan produk peralatan/elektronik (yang memerlukan jaminan masa garansi). Apabila entitas-entitas tersebut dibuat sebagai class independen, akan terjadi redundansi atau duplikasi atribut umum (kode, nama, harga, stok) dan method mutator yang identik. Oleh karena itu, pada praktikum P5 ini diterapkan prinsip Generalization dan Inheritance untuk membentuk arsitektur Superclass dan Subclass yang terstruktur.", bold_prefix="1.2 Latar Belakang Masalah: ")

    add_styled_heading(doc, "2. Perencanaan Agile Sprint (Sprint P5)", level=1)
    add_body_p(doc, "Memperbaiki desain proyek Sistem Kasir Sederhana dengan mengidentifikasi class yang memiliki karakteristik serupa, melakukan generalization untuk membentuk Superclass Produk, serta mengimplementasikan Subclass ProdukMakanan dan ProdukElektronik menggunakan klausa extends dan super, dengan tetap mempertahankan seluruh relasi P4.", bold_prefix="2.1 Sprint Goal: ")

    add_body_p(doc, "Sprint Backlog P5", bold_prefix="2.2 ")
    t_sb = doc.add_table(rows=8, cols=3)
    t_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_sb)
    sb_data = [
        ("ID", "Sprint Backlog", "Status"),
        ("SB-01", "Mencari dan mengaudit class yang memiliki potensi duplikasi atribut/behavior.", "DONE"),
        ("SB-02", "Menentukan dan mendefinisikan Superclass (Produk).", "DONE"),
        ("SB-03", "Menentukan dan mendefinisikan Subclass (ProdukMakanan dan ProdukElektronik).", "DONE"),
        ("SB-04", "Memperbarui UML Class Diagram proyek sebelum dan sesudah refactoring.", "DONE"),
        ("SB-05", "Implementasi pewarisan menggunakan kata kunci extends pada subclass.", "DONE"),
        ("SB-06", "Implementasi pemanggilan constructor superclass menggunakan super(...).", "DONE"),
        ("SB-07", "Melakukan pengujian menyeluruh pada Main.java mencakup transaksi multi-kategori.", "DONE")
    ]
    for r_idx, row_content in enumerate(sb_data):
        for c_idx, val in enumerate(row_content):
            cell = t_sb.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(0.9)
            elif c_idx == 1:
                cell.width = Inches(4.7)
            else:
                cell.width = Inches(0.9)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0 or c_idx == 2:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "3. Audit Class dan Identifikasi Generalization (Bagian A)", level=1)
    add_body_p(doc, "Berdasarkan hasil analisis, ditemukan entitas produk makanan dan produk elektronik yang memiliki karakteristik umum identik. Hasil audit dikelompokkan ke dalam tabel berikut:")

    t_audit = doc.add_table(rows=3, cols=3)
    t_audit.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_audit)
    audit_data = [
        ("Class", "Attribute Umum (Kandidat Superclass)", "Attribute Khusus (Subclass)"),
        ("ProdukMakanan", "kode, nama, harga, stok", "tanggalKadaluarsa (String)"),
        ("ProdukElektronik", "kode, nama, harga, stok", "garansiBulan (int)")
    ]
    for r_idx, row_content in enumerate(audit_data):
        for c_idx, val in enumerate(row_content):
            cell = t_audit.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.8)
            elif c_idx == 1:
                cell.width = Inches(2.7)
            else:
                cell.width = Inches(2.0)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_body_p(doc, "Penetapan Struktur Hierarki:", bold_prefix="3.1 ")
    add_bullet_p(doc, "Merupakan kelas induk yang memegang atribut bersama: kode, nama, harga, stok, serta method operasional (kurangiStok, tambahStok, hitungNilaiInventaris).", bold_prefix="• Superclass: Produk — ")
    add_bullet_p(doc, "Mewarisi Produk dengan penambahan atribut khusus tanggalKadaluarsa.", bold_prefix="• Subclass 1: ProdukMakanan — ")
    add_bullet_p(doc, "Mewarisi Produk dengan penambahan atribut khusus garansiBulan.", bold_prefix="• Subclass 2: ProdukElektronik — ")

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 2 / BAGIAN 2: CLASS DIAGRAM SEBELUM/SESUDAH & HUBUNGAN IS-A
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "4. Pembaruan Class Diagram (Bagian B)", level=1)
    
    add_body_p(doc, "Class Diagram Sebelum Refactoring (P4)", bold_prefix="4.1 ")
    add_body_p(doc, "Sebelum refactoring, jika class ProdukMakanan dan ProdukElektronik didefinisikan secara mandiri, terjadi duplikasi atribut kode, nama, harga, stok secara terulang pada masing-masing kelas.")

    path_sebelum = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\assets", "class_diagram_sebelum.png")
    if os.path.exists(path_sebelum):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.space_before = Pt(2)
        p_img1.paragraph_format.space_after = Pt(2)
        run_img1 = p_img1.add_run()
        run_img1.add_picture(path_sebelum, width=Inches(5.8))

        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1.paragraph_format.space_before = Pt(0)
        p_cap1.paragraph_format.space_after = Pt(6)
        r_cap1 = p_cap1.add_run("Gambar 1. UML Class Diagram Sebelum Refactoring (Terjadi Duplikasi Atribut)")
        r_cap1.font.name = "Times New Roman"
        r_cap1.font.size = Pt(9.5)
        r_cap1.font.italic = True
        r_cap1.font.color.rgb = RGBColor(0, 0, 0)

    add_body_p(doc, "Class Diagram Setelah Refactoring (P5: Dengan Inheritance)", bold_prefix="4.2 ")
    add_body_p(doc, "Setelah refactoring, atribut umum diekstrak ke dalam Superclass Produk. Subclass ProdukMakanan dan ProdukElektronik mewarisi Produk menggunakan simbol segitiga generalization (▲), sementara relasi P4 (Association dengan Pelanggan, Aggregation dengan ItemTransaksi, dan Composition dengan Transaksi) tetap dipertahankan utuh.")

    path_sesudah = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\assets", "class_diagram_sesudah.png")
    if os.path.exists(path_sesudah):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(2)
        p_img2.paragraph_format.space_after = Pt(2)
        run_img2 = p_img2.add_run()
        run_img2.add_picture(path_sesudah, width=Inches(6.2))

        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_before = Pt(0)
        p_cap2.paragraph_format.space_after = Pt(6)
        r_cap2 = p_cap2.add_run("Gambar 2. UML Class Diagram Setelah Refactoring (Penerapan Inheritance & Generalization P5)")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(9.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(0, 0, 0)

    add_styled_heading(doc, "5. Penjelasan Hubungan is-a dan Manfaat Refactoring", level=1)
    add_body_p(doc, "Inheritance pada perancangan ini memenuhi prinsip hubungan 'is-a' secara sah dan logis:", bold_prefix="5.1 Pemenuhan Prinsip is-a: ")
    add_bullet_p(doc, "ProdukMakanan is-a Produk (Produk makanan merupakan jenis produk yang sah dan memiliki kode, nama, harga, stok, serta karakteristik kadaluarsa).", bold_prefix="• ")
    add_bullet_p(doc, "ProdukElektronik is-a Produk (Produk elektronik merupakan jenis produk yang sah dan memiliki kode, nama, harga, stok, serta jaminan masa garansi).", bold_prefix="• ")
    
    add_body_p(doc, "Manfaat Nyata Refactoring:", bold_prefix="5.2 ")
    add_bullet_p(doc, "Atribut kode, nama, harga, stok serta logika validasinya cukup didefinisikan satu kali pada Superclass Produk.", bold_prefix="1. Mengeliminasi Duplikasi Kode: ")
    add_bullet_p(doc, "Pemisahan antara karakteristik umum dan karakteristik khusus membuat model data lebih intuitif dan rapi.", bold_prefix="2. Struktur Class Lebih Bersih: ")
    add_bullet_p(doc, "Class ItemTransaksi dan Transaksi dari P4 tidak perlu diubah karena keduanya berinteraksi dengan Superclass Produk secara polimorfik.", bold_prefix="3. Kemudahan Maintenance & Reusability: ")

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 3 / BAGIAN 3: PENGUJIAN, ANALISIS EXTENDS & SUPER, SPRINT REVIEW & RETROSPECTIVE
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "6. Pengujian Program (Bagian F) dan Hasil Eksekusi", level=1)
    add_body_p(doc, "Pengujian pada class Main.java membuktikan instansiasi objek dari kedua subclass, pemanggilan constructor superclass menggunakan super(...), pewarisan atribut/method, serta integrasi transaksi belanja multi-kategori (makanan dan elektronik dalam satu struk).")

    add_body_p(doc, "Tangkapan Layar Hasil Running Terminal (Monokrom)", bold_prefix="6.1 ")
    path_term = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\assets", "hasil_running.png")
    if os.path.exists(path_term):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.paragraph_format.space_before = Pt(2)
        p_img3.paragraph_format.space_after = Pt(2)
        run_img3 = p_img3.add_run()
        run_img3.add_picture(path_term, width=Inches(5.9))

        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap3.paragraph_format.space_before = Pt(0)
        p_cap3.paragraph_format.space_after = Pt(6)
        r_cap3 = p_cap3.add_run("Gambar 3. Bukti Running Program Main.java pada Terminal Windows PowerShell")
        r_cap3.font.name = "Times New Roman"
        r_cap3.font.size = Pt(9.5)
        r_cap3.font.italic = True
        r_cap3.font.color.rgb = RGBColor(0, 0, 0)

    add_styled_heading(doc, "7. Bukti dan Analisis Penggunaan extends dan super (Bagian E)", level=1)
    add_body_p(doc, "Subclass mendeklarasikan hubungan inheritance melalui klausa 'public class ProdukMakanan extends Produk' dan 'public class ProdukElektronik extends Produk'. Hal ini memungkinkan kedua subclass mewarisi method publik dari superclass seperti getNama(), getHarga(), kurangiStok(), dan hitungNilaiInventaris().", bold_prefix="7.1 Penerapan extends: ")
    add_body_p(doc, "Pada constructor kedua subclass, pemanggilan 'super(kode, nama, harga, stok);' diletakkan pada baris pertama. Pemanggilan ini meneruskan argumen inisialisasi ke constructor Superclass Produk. Urutan eksekusi constructor berjalan secara hierarkis: constructor Produk dijalankan terlebih dahulu untuk menginisialisasi atribut dasar beserta validasinya, kemudian dilanjutkan eksekusi constructor subclass untuk menginisialisasi atribut spesifik (tanggalKadaluarsa atau garansiBulan).", bold_prefix="7.2 Penerapan super(...) & Urutan Eksekusi: ")

    add_styled_heading(doc, "8. Sprint Review (Evaluasi Sprint P5)", level=1)
    t_rev = doc.add_table(rows=9, cols=2)
    t_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_rev)
    rev_data = [
        ("Item Evaluasi", "Hasil / Keterangan"),
        ("Kandidat inheritance ditemukan", "Ditemukan entitas ProdukMakanan dan ProdukElektronik dengan atribut umum yang sama."),
        ("Superclass berhasil dibuat", "Superclass Produk berhasil didefinisikan dengan atribut umum dan method terenkapsulasi."),
        ("Minimal 2 subclass dibuat", "Dibuat 2 subclass: ProdukMakanan dan ProdukElektronik."),
        ("extends diterapkan", "Klausa extends diterapkan pada kedua subclass untuk mewarisi class Produk."),
        ("super diterapkan", "Kata kunci super(...) diterapkan pada constructor subclass untuk inisialisasi superclass."),
        ("Class diagram diperbarui", "Class diagram diperbarui dengan simbol generalization dan perbandingan sebelum/sesudah."),
        ("Program berhasil dijalankan", "Program sukses dikompilasi (javac) dan dieksekusi (java) 100% bebas error."),
        ("Kendala", "Tidak ditemukan kendala fatal. Seluruh relasi P4 terintegrasi mulus dengan inheritance P5.")
    ]
    for r_idx, row_content in enumerate(rev_data):
        for c_idx, val in enumerate(row_content):
            cell = t_rev.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(2.3)
            else:
                cell.width = Inches(4.2)
            set_cell_margins(cell, 50, 50, 70, 70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "9. Sprint Retrospective", level=1)
    add_body_p(doc, "Penerapan konsep generalization dan inheritance berhasil mengeliminasi duplikasi kode pada model produk. Pemanggilan super(...) menjamin inisialisasi bertingkat yang kokoh, dan integrasi multi-item belanja berjalan sangat stabil.", bold_prefix="What Went Well? ")
    add_body_p(doc, "Perlu kecermatan dalam memastikan access modifier superclass tetap menjaga prinsip information hiding tanpa membatasi fleksibilitas subclass.", bold_prefix="What Went Wrong? ")
    add_body_p(doc, "Pada sprint berikutnya (P6: Polymorphism), behavior perhitungan diskon atau tampilan produk antar-subclass akan dikembangkan melalui method overriding dinamis.", bold_prefix="Improvement: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "10. Definition of Done (DoD) Checklist Modul 5", level=1)
    dod_items = [
        ("Menggunakan project P4 (Sistem Kasir Sederhana)", True),
        ("Melakukan audit terhadap class dan menemukan duplikasi", True),
        ("Menemukan kandidat generalization yang relevan (Produk)", True),
        ("Minimal 1 superclass dibuat (Produk)", True),
        ("Minimal 2 subclass dibuat (ProdukMakanan dan ProdukElektronik)", True),
        ("Menggunakan kata kunci extends pada subclass", True),
        ("Menggunakan super(...) pada constructor subclass", True),
        ("Constructor superclass dan subclass berjalan sesuai urutan", True),
        ("Tidak terdapat duplikasi attribute umum pada subclass", True),
        ("Hubungan inheritance memenuhi prinsip is-a", True),
        ("Class Diagram diperbarui (perbandingan sebelum dan sesudah)", True),
        ("Relasi dari P4 tetap dipertahankan secara utuh", True),
        ("Program berhasil dikompilasi bebas error", True),
        ("Program berhasil dijalankan dengan output yang sesuai", True),
        ("Sprint Backlog diperbarui dengan status DONE", True),
        ("Sprint Review dan Sprint Retrospective tersedia lengkap", True)
    ]
    t_dod = doc.add_table(rows=len(dod_items)+1, cols=2)
    t_dod.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_dod)
    
    hdr0, hdr1 = t_dod.rows[0].cells[0], t_dod.rows[0].cells[1]
    hdr0.width, hdr1.width = Inches(5.3), Inches(1.2)
    set_cell_margins(hdr0, 60, 60, 80, 80)
    set_cell_margins(hdr1, 60, 60, 80, 80)
    p = hdr0.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Kriteria Definition of Done (Modul 5)"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 0, 0)
    p = hdr1.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Status"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 0, 0)

    for idx, (item, fulfilled) in enumerate(dod_items, start=1):
        c0, c1 = t_dod.rows[idx].cells[0], t_dod.rows[idx].cells[1]
        c0.width, c1.width = Inches(5.3), Inches(1.2)
        set_cell_margins(c0, 40, 40, 70, 70)
        set_cell_margins(c1, 40, 40, 70, 70)
        
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0); p0.paragraph_format.line_spacing = 1.1
        r0 = p0.add_run(item); r0.font.name = "Times New Roman"; r0.font.size = Pt(9); r0.font.color.rgb = RGBColor(0, 0, 0)
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0); p1.paragraph_format.line_spacing = 1.1
        r1 = p1.add_run("[ v ] Terpenuhi" if fulfilled else "[   ] Belum"); r1.font.name = "Times New Roman"; r1.font.size = Pt(9); r1.font.bold = True; r1.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # LAMPIRAN: SOURCE CODE JAVA LENGKAP P5 (Hitam & Putih)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "Lampiran: Source Code Java (Inheritance P5)", level=1)
    
    code_files = [
        ("Produk.java (Superclass)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\Produk.java"),
        ("ProdukMakanan.java (Subclass 1)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\ProdukMakanan.java"),
        ("ProdukElektronik.java (Subclass 2)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\ProdukElektronik.java"),
        ("Pelanggan.java (Class Relasi P4)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\Pelanggan.java"),
        ("ItemTransaksi.java (Class Relasi P4)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\ItemTransaksi.java"),
        ("Transaksi.java (Class Relasi P4)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\Transaksi.java"),
        ("Main.java (Main Runner Test)", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\src\Main.java")
    ]

    for fname, fpath in code_files:
        add_styled_heading(doc, f"Berkas: {fname}", level=2)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                code_content = f.read()
            
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(6.5)
            set_cell_margins(cell, 80, 80, 100, 100)
            set_table_borders_bw(tbl)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(code_content)
            r.font.name = "Courier New"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0, 0, 0)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_filename)
    print(f"P5 Document saved successfully: {output_filename}")

out_docx = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\Laporan_P5_3125522007_Rafael_Rizky.docx"
build_laporan_p5(out_docx)

readme_docx = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P5_3125522007_Rafael_Rizky\README.docx"
build_laporan_p5(readme_docx)

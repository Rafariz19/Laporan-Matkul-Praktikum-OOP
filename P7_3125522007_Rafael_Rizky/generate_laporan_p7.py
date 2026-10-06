import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P7_3125522007_Rafael_Rizky"
assets_dir = os.path.join(base_dir, "assets")

def set_cell_margins(cell, top=90, bottom=90, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_bw(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11.5)
    r_body.font.italic = italic
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_code_block_p(doc, code_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str)
    run.font.name = "Courier New"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_laporan_p7(output_filename):
    doc = docx.Document()

    # Margin 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # HEADER / TITLE (Pure Black & White)
    # ---------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("LAPORAN PRAKTIKUM PEMROGRAMAN BERORIENTASI OBYEK")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r = p_sub.add_run("MODUL 7: ABSTRACT CLASS, ABSTRACT METHOD, DAN INTERFACE")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    # Identity Table (Pure Black & White)
    table_id = doc.add_table(rows=4, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False
    
    id_data = [
        ("Nama Mahasiswa / NRP", ": Rafael Rizky / 3125522007"),
        ("Dosen Pengampu", ": Nirwana Haidar Hari, S.Pd., M.Kom."),
        ("Institusi & Program", ": PENS PSDKU Sumenep (Tahun 2026)"),
        ("Judul Proyek Lanjutan", ": Sistem Kasir Sederhana (Sprint P7 - Refactoring Abstraction & Interface)")
    ]

    for row_idx, (col0, col1) in enumerate(id_data):
        row = table_id.rows[row_idx]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        for c_idx, text in enumerate([col0, col1]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if c_idx == 0:
                run.font.bold = True

    set_table_borders_bw(table_id)

    # Divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(10)
    p_div.paragraph_format.space_after = Pt(10)
    r = p_div.add_run("_________________________________________________________________________________")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # BAB I: AGILE SPRINT - P7
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "I. AGILE SPRINT PLANNING — SPRINT P7", level=1)
    
    add_body_p(doc, 
               "Pada penugasan Praktikum Modul 7, dilakukan pengembangan desain berbasis Agile Sprint terhadap proyek kelanjutan P6 "
               "bernama Sistem Kasir Sederhana. Refactoring difokuskan pada penguatan arsitektur berorientasi obyek menggunakan prinsip Abstraction "
               "(Abstract Class dan Abstract Method) serta kontrak antarmuka (Interface), guna memastikan pemisahan konsep umum dan perilaku spesifik "
               "berjalan konsisten tanpa mengorbankan stabilitas fitur bisnis yang telah dibangun dari P1 hingga P6.")

    add_body_p(doc, 
               "Mengembangkan desain proyek P6 dengan menerapkan abstract class dan interface yang sesuai sehingga struktur class lebih jelas, "
               "perilaku object memiliki kontrak yang konsisten, dan seluruh alur program kasir tetap dapat dijalankan secara sempurna.",
               bold_prefix="Sprint Goal P7: ")

    add_styled_heading(doc, "Tabel 1. Sprint Backlog Penugasan P7", level=2)
    tbl_sb = doc.add_table(rows=8, cols=3)
    tbl_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sb.autofit = False

    sb_headers = ["ID Backlog", "Deskripsi Aktivitas Sprint Backlog P7", "Status Pengerjaan"]
    for i, h in enumerate(sb_headers):
        cell = tbl_sb.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    sb_data = [
        ("SB-01", "Audit hierarchy dan behavior proyek P6 (evaluasi instansiasi Superclass Produk)", "DONE (100%)"),
        ("SB-02", "Menentukan kandidat abstract class (transformasi class Produk menjadi abstract)", "DONE (100%)"),
        ("SB-03", "Menentukan abstract method (getKategoriInfo() dan tampilkanDetailKhusus())", "DONE (100%)"),
        ("SB-04", "Menentukan kandidat interface (pembuatan kontrak DapatDidiskon)", "DONE (100%)"),
        ("SB-05", "Memperbarui class diagram dengan notasi generalization dan realization", "DONE (100%)"),
        ("SB-06", "Mengimplementasikan abstract class Produk dan interface DapatDidiskon pada subclass", "DONE (100%)"),
        ("SB-07", "Menguji polymorphism melalui 5 skenario komprehensif dan verifikasi program kasir", "DONE (100%)")
    ]

    widths = [Inches(1.2), Inches(3.8), Inches(1.5)]
    for r_idx, (col0, col1, col2) in enumerate(sb_data, start=1):
        row = tbl_sb.rows[r_idx]
        for c_idx, val in enumerate([col0, col1, col2]):
            cell = row.cells[c_idx]
            cell.width = widths[c_idx]
            p = cell.paragraphs[0]
            if c_idx in [0, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if c_idx == 2:
                run.font.bold = True
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    set_table_borders_bw(tbl_sb)

    # ---------------------------------------------------------------------------
    # BAB II: AUDIT DESAIN P6 & KEPUTUSAN ABSTRACTION (BAGIAN A)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "II. AUDIT DESAIN P6 DAN KEPUTUSAN REFACTORING (BAGIAN A)", level=1)
    
    add_body_p(doc, 
               "Pada evaluasi arsitektur P6, class Produk diperlakukan sebagai superclass konkret yang dapat diinstansiasi secara bebas menggunakan keyword new. "
               "Dalam analisis bisnis sistem kasir riil, kondisi ini memicu celah rancangan (design flaw) karena kasir tidak pernah menjual barang generik "
               "tanpa karakteristik kategori konkret (seperti tanggal kadaluarsa pangan atau masa garansi perangkat keras). Oleh karena itu, diputuskan "
               "sejumlah perubahan struktural berlandaskan prinsip OOP murni:")

    add_bullet_p(doc, "Diubah menjadi public abstract class. Class ini memegang tanggung jawab enkapsulasi atribut umum (kode, nama, harga, stok) "
                      "dan operasi inventaris, namun dilarang keras diinstansiasi secara langsung.", bold_prefix="Superclass Produk: ")
    add_bullet_p(doc, "Dideklarasikan sebagai abstract method (public abstract String getKategoriInfo() dan public abstract void tampilkanDetailKhusus()). "
                      "Setiap subclass konkret diwajibkan menyediakan realisasi perilakunya masing-masing.", bold_prefix="Metode Konseptual: ")
    add_bullet_p(doc, "Dibuat kontrak antarmuka baru public interface DapatDidiskon yang menaungi method hitungDiskon(double) dan getHargaSetelahDiskon(double), "
                      "memungkinkan objek-objek komoditas toko memenuhi kontrak promosi diskon barang.", bold_prefix="Interface Komoditas: ")

    add_styled_heading(doc, "Tabel 2. Matriks Keputusan Perubahan Desain dari P6 ke P7", level=2)
    tbl_audit = doc.add_table(rows=6, cols=4)
    tbl_audit.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_audit.autofit = False

    audit_headers = ["Class / Behavior", "Kondisi P6", "Rencana Desain P7", "Alasan Keputusan Desain"]
    for i, h in enumerate(audit_headers):
        cell = tbl_audit.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)

    audit_rows = [
        ("Class Produk", "Class konkret biasa (dapat diinstansiasi)", "Abstract Class (public abstract class)",
         "Hanya mewakili konsep umum komoditas toko; instansiasi objek generik dilarang demi integritas domain kasir."),
        ("getKategoriInfo()", "Method biasa mengembalikan string standar", "Abstract Method (tanpa body di superclass)",
         "Setiap komoditas wajib mengklasifikasikan kategorinya secara eksplisit dan mandiri."),
        ("tampilkanDetailKhusus()", "Belum ada (digabung pada overriding)", "Abstract Method (kontrak tampilan khusus)",
         "Memaksa subclass menampilkan atribut spesifiknya (kadaluarsa/garansi) lewat Template Method pada Produk."),
        ("DapatDidiskon", "Belum tersedia", "Interface (kontrak perhitungan promo)",
         "Menyediakan standardisasi kontrak diskon promosi barang tanpa membatasi hierarki pewarisan class."),
        ("ProdukMakanan &\nProdukElektronik", "Subclass biasa (extends Produk)", "Subclass konkret (extends Produk implements DapatDidiskon)",
         "Memenuhi kewajiban abstract method superclass sekaligus merealisasikan kontrak promosi diskon interface.")
    ]

    a_widths = [Inches(1.4), Inches(1.5), Inches(1.7), Inches(1.9)]
    for r_idx, (c0, c1, c2, c3) in enumerate(audit_rows, start=1):
        row = tbl_audit.rows[r_idx]
        for c_idx, val in enumerate([c0, c1, c2, c3]):
            cell = row.cells[c_idx]
            cell.width = a_widths[c_idx]
            p = cell.paragraphs[0]
            if c_idx in [0, 1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if c_idx == 0:
                run.font.bold = True
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    set_table_borders_bw(tbl_audit)

    add_styled_heading(doc, "Tabel 3. Komparasi Karakteristik Abstract Class dan Interface", level=2)
    tbl_diff = doc.add_table(rows=6, cols=3)
    tbl_diff.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_diff.autofit = False

    diff_headers = ["Karakteristik OOP", "Abstract Class (Produk)", "Interface (DapatDidiskon)"]
    for i, h in enumerate(diff_headers):
        cell = tbl_diff.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_margins(cell, top=70, bottom=70, left=90, right=90)

    diff_rows = [
        ("Keyword Definisi & Pewarisan", "abstract class dan extends", "interface dan implements"),
        ("Attribute Instance", "Dapat memiliki attribute instance (kode, nama, harga, stok)", "Tidak dapat memiliki attribute instance"),
        ("Constructor", "Memiliki constructor untuk inisialisasi state superclass", "Tidak memiliki constructor"),
        ("Realisasi Method", "Dapat memiliki method konkret ber-body dan abstract method", "Secara default seluruh method adalah public abstract"),
        ("Dukungan Hubungan", "Single inheritance (hanya satu superclass langsung)", "Multiple interface implementation (dapat implementasi banyak)")
    ]

    d_widths = [Inches(1.8), Inches(2.4), Inches(2.3)]
    for r_idx, (c0, c1, c2) in enumerate(diff_rows, start=1):
        row = tbl_diff.rows[r_idx]
        for c_idx, val in enumerate([c0, c1, c2]):
            cell = row.cells[c_idx]
            cell.width = d_widths[c_idx]
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if c_idx == 0:
                run.font.bold = True
            set_cell_margins(cell, top=55, bottom=55, left=90, right=90)

    set_table_borders_bw(tbl_diff)

    # ---------------------------------------------------------------------------
    # BAB III: UML CLASS DIAGRAM TERBARU (BAGIAN B)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "III. PERUBAHAN UML CLASS DIAGRAM P7 (BAGIAN B)", level=1)
    
    add_body_p(doc, 
               "Class Diagram sistem kasir telah disempurnakan dengan mengintegrasikan konsep Abstraction dan Realization sesuai standar UML formal:")
    
    add_bullet_p(doc, "Dilambangkan dengan garis lurus berujung segitiga berongga tertutup (▲) yang mengarah dari subclass ProdukMakanan dan ProdukElektronik menuju abstract class Produk.", bold_prefix="Generalization (Inheritance): ")
    add_bullet_p(doc, "Dilambangkan dengan garis putus-putus berujung segitiga berongga tertutup (┆▲) yang mengarah dari ProdukMakanan dan ProdukElektronik menuju <<interface>> DapatDidiskon.", bold_prefix="Realization (Interface Contract): ")
    add_bullet_p(doc, "Transaksi terhubung ke Pelanggan dengan garis berpanah terbuka (-->).", bold_prefix="Association (uses-a): ")
    add_bullet_p(doc, "ItemTransaksi memiliki referensi polymorphic ke tipe abstract class Produk dilambangkan belah ketupat putih (◇-->).", bold_prefix="Aggregation (has-a): ")
    add_bullet_p(doc, "Transaksi memiliki siklus hidup yang mengontrol kumpulan ItemTransaksi dilambangkan belah ketupat hitam (◆-->).", bold_prefix="Composition (part-of): ")

    class_diagram_path = os.path.join(assets_dir, "class_diagram_p7.png")
    if os.path.exists(class_diagram_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(class_diagram_path, width=Inches(6.3))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Gambar 1. UML Class Diagram Refactoring Abstraction dan Realization (P7)")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # BAB IV: IMPLEMENTASI SOURCE CODE JAVA (BAGIAN C, D, E)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "IV. IMPLEMENTASI SOURCE CODE JAVA (BAGIAN C, D, E)", level=1)

    add_styled_heading(doc, "4.1 Abstract Class: Produk.java (Bagian C)", level=2)
    add_body_p(doc, 
               "Class Produk dideklarasikan secara abstract sehingga tidak dapat diinstansiasi dengan operator new. "
               "Class ini menyediakan attribute dasar, constructor berenkapsulasi, operasi stok, dua abstract method, "
               "dan satu template method konkret tampilkanData() yang mendistribusikan pemanggilan secara dinamis ke subclass.")
    
    code_produk = """public abstract class Produk {
    private String kode;
    private String nama;
    private int harga;
    private int stok;

    public Produk(String kode, String nama, int harga, int stok) {
        this.kode = (kode != null && !kode.trim().isEmpty()) ? kode : "PROD-DEF";
        setNama(nama);
        setHarga(harga);
        setStok(stok);
    }

    // Getter dan Setter dengan Validasi Ketat
    public String getKode() { return kode; }
    public String getNama() { return nama; }
    public int getHarga() { return harga; }
    public int getStok() { return stok; }
    public void setNama(String nama) { if (nama != null && !nama.trim().isEmpty()) this.nama = nama; }
    public void setHarga(int harga) { if (harga > 0) this.harga = harga; }
    public void setStok(int stok) { if (stok >= 0) this.stok = stok; }

    public void tambahStok(int jumlah) { if (jumlah > 0) this.stok += jumlah; }
    public boolean kurangiStok(int jumlah) {
        if (jumlah > 0 && this.stok >= jumlah) { this.stok -= jumlah; return true; }
        return false;
    }
    public int hitungNilaiInventaris() { return this.harga * this.stok; }

    // ABSTRACT METHODS (Wajib direalisasikan Subclass Konkret)
    public abstract String getKategoriInfo();
    public abstract void tampilkanDetailKhusus();

    // TEMPLATE METHOD KONKRET (Memanfaatkan Dynamic Binding)
    public void tampilkanData() {
        System.out.println("--------------------------------------------------");
        System.out.println("Kategori    : " + getKategoriInfo());
        System.out.println("Kode        : " + getKode());
        System.out.println("Nama        : " + getNama());
        System.out.println("Harga       : Rp " + String.format("%,d", getHarga()).replace(',', '.'));
        System.out.println("Stok        : " + getStok() + " unit");
        System.out.println("Inventaris  : Rp " + String.format("%,d", hitungNilaiInventaris()).replace(',', '.'));
        tampilkanDetailKhusus();
        System.out.println("--------------------------------------------------");
    }
}"""
    add_code_block_p(doc, code_produk)

    add_styled_heading(doc, "4.2 Interface: DapatDidiskon.java (Bagian E)", level=2)
    add_body_p(doc, 
               "Interface DapatDidiskon bertindak sebagai kontrak bisnis yang menetapkan kewajiban setiap produk berdiskon "
               "untuk mampu menghitung potongan nominal dan menentukan harga akhir setelah diskon promosi.")
    
    code_interface = """public interface DapatDidiskon {
    // Kontrak Interface Perhitungan Diskon Promosi
    double hitungDiskon(double persentase);
    double getHargaSetelahDiskon(double persentase);
}"""
    add_code_block_p(doc, code_interface)

    add_styled_heading(doc, "4.3 Subclass 1: ProdukMakanan.java (Bagian D & E)", level=2)
    add_body_p(doc, 
               "ProdukMakanan mewarisi Produk dan mengimplementasikan DapatDidiskon. Subclass ini menambahkan attribute tanggalKadaluarsa, "
               "merealisasikan abstract method getKategoriInfo() & tampilkanDetailKhusus(), serta merealisasikan kontrak interface DapatDidiskon.")

    code_makanan = """public class ProdukMakanan extends Produk implements DapatDidiskon {
    private String tanggalKadaluarsa;

    public ProdukMakanan(String kode, String nama, int harga, int stok, String tanggalKadaluarsa) {
        super(kode, nama, harga, stok);
        setTanggalKadaluarsa(tanggalKadaluarsa);
    }

    public String getTanggalKadaluarsa() { return tanggalKadaluarsa; }
    public void setTanggalKadaluarsa(String tgl) {
        this.tanggalKadaluarsa = (tgl != null && !tgl.trim().isEmpty()) ? tgl : "01-01-2027";
    }

    // Realisasi Abstract Method dari Superclass Produk
    @Override
    public String getKategoriInfo() { return "Makanan & Minuman Segar (Konsumsi)"; }

    @Override
    public void tampilkanDetailKhusus() { System.out.println("Kadaluarsa  : " + tanggalKadaluarsa); }

    // Realisasi Kontrak Interface DapatDidiskon
    @Override
    public double hitungDiskon(double persentase) {
        if (persentase <= 0.0) return 0.0;
        if (persentase > 100.0) persentase = 100.0;
        return getHarga() * (persentase / 100.0);
    }

    @Override
    public double getHargaSetelahDiskon(double persentase) {
        return getHarga() - hitungDiskon(persentase);
    }
}"""
    add_code_block_p(doc, code_makanan)

    add_styled_heading(doc, "4.4 Subclass 2: ProdukElektronik.java (Bagian D & E)", level=2)
    add_body_p(doc, 
               "ProdukElektronik mewarisi Produk dan mengimplementasikan DapatDidiskon. Subclass ini menambahkan attribute garansiBulan, "
               "merealisasikan abstract method getKategoriInfo() & tampilkanDetailKhusus(), serta merealisasikan kontrak interface DapatDidiskon.")

    code_elektro = """public class ProdukElektronik extends Produk implements DapatDidiskon {
    private int garansiBulan;

    public ProdukElektronik(String kode, String nama, int harga, int stok, int garansiBulan) {
        super(kode, nama, harga, stok);
        setGaransiBulan(garansiBulan);
    }

    public int getGaransiBulan() { return garansiBulan; }
    public void setGaransiBulan(int garansiBulan) {
        this.garansiBulan = (garansiBulan >= 0) ? garansiBulan : 0;
    }

    // Realisasi Abstract Method dari Superclass Produk
    @Override
    public String getKategoriInfo() { return "Elektronik & Aksesoris Gadget (Hardware)"; }

    @Override
    public void tampilkanDetailKhusus() { System.out.println("Garansi     : " + garansiBulan + " Bulan (Garansi Resmi)"); }

    // Realisasi Kontrak Interface DapatDidiskon
    @Override
    public double hitungDiskon(double persentase) {
        if (persentase <= 0.0) return 0.0;
        if (persentase > 100.0) persentase = 100.0;
        return getHarga() * (persentase / 100.0);
    }

    @Override
    public double getHargaSetelahDiskon(double persentase) {
        return getHarga() - hitungDiskon(persentase);
    }
}"""
    add_code_block_p(doc, code_elektro)

    # ---------------------------------------------------------------------------
    # BAB V: PENGUJIAN POLYMORPHISM & VERIFIKASI (BAGIAN F)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "V. PENGUJIAN POLYMORPHISM DAN PROGRAM KASIR (BAGIAN F)", level=1)
    
    add_body_p(doc, 
               "Untuk memenuhi persyaratan Modul 7 Bagian F (minimal 4 skenario), dirancang 5 skenario pengujian menyeluruh pada file Main.java:")

    add_bullet_p(doc, "Membuktikan instansiasi objek konkret ProdukMakanan dan ProdukElektronik, serta mengonfirmasi bahwa compiler menolak new Produk(...) karena class bersifat abstract.", bold_prefix="Skenario 1 (Instansiasi Subclass & Abstraction): ")
    add_bullet_p(doc, "Membuktikan upcasting objek subclass ke variabel referensi Produk. Dynamic binding memanggil method overriding getKategoriInfo() dan tampilkanData() milik objek aktual saat runtime.", bold_prefix="Skenario 2 (Abstract Method via Superclass Ref): ")
    add_bullet_p(doc, "Membuktikan polymorphic dispatch melalui tipe referensi interface DapatDidiskon (hitungDiskon() dan getHargaSetelahDiskon()).", bold_prefix="Skenario 3 (Interface Method via Interface Ref): ")
    add_bullet_p(doc, "Membuktikan pemrosesan massal menggunakan array Produk[] dan array DapatDidiskon[] dalam simulasi flash sale diskon 20%.", bold_prefix="Skenario 4 (Polymorphic Collections): ")
    add_bullet_p(doc, "Membuktikan seluruh relasi P4-P6 (Association, Aggregation, Composition, Overloading) beroperasi secara harmonis tanpa regresi.", bold_prefix="Skenario 5 (Integrasi Transaksi Kasir): ")

    running_img_path = os.path.join(assets_dir, "hasil_running.png")
    if os.path.exists(running_img_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(8)
        p_img2.paragraph_format.space_after = Pt(2)
        doc.add_picture(running_img_path, width=Inches(6.3))

        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_before = Pt(2)
        p_cap2.paragraph_format.space_after = Pt(12)
        r_cap2 = p_cap2.add_run("Gambar 2. Bukti Running Program Main.java pada Terminal (Eksekusi 5 Skenario P7)")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(0, 0, 0)

    add_styled_heading(doc, "Tabel 4. Matriks Hasil Pengujian Skenario Polymorphism P7", level=2)
    tbl_scen = doc.add_table(rows=6, cols=4)
    tbl_scen.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_scen.autofit = False

    scen_headers = ["No.", "Skenario Pengujian", "Hasil yang Diharapkan", "Status Verifikasi"]
    for i, h in enumerate(scen_headers):
        cell = tbl_scen.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)

    scen_rows = [
        ("1", "Instansiasi objek subclass konkret", "Subclass berhasil diinstansiasi; new Produk(...) ditolak compiler", "SUKSES (100%)"),
        ("2", "Memanggil abstract method via reference superclass Produk", "JVM mengeksekusi method konkret subclass melalui dynamic binding", "SUKSES (100%)"),
        ("3", "Memanggil method via reference interface DapatDidiskon", "Kalkulasi diskon berjalan akurat sesuai objek aktual yang diacu", "SUKSES (100%)"),
        ("4", "Mengiterasi polymorphic collection (array Produk[] & DapatDidiskon[])", "Setiap objek menjalankan perilakunya masing-masing dalam satu perulangan", "SUKSES (100%)"),
        ("5", "Integrasi transaksi penjualan kasir dan pembayaran tunai", "Struk tercetak rapi, kalkulasi kembalian tunai akurat, seluruh fitur P1-P6 utuh", "SUKSES (100%)")
    ]

    s_widths = [Inches(0.6), Inches(2.2), Inches(2.5), Inches(1.2)]
    for r_idx, (c0, c1, c2, c3) in enumerate(scen_rows, start=1):
        row = tbl_scen.rows[r_idx]
        for c_idx, val in enumerate([c0, c1, c2, c3]):
            cell = row.cells[c_idx]
            cell.width = s_widths[c_idx]
            p = cell.paragraphs[0]
            if c_idx in [0, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if c_idx == 3:
                run.font.bold = True
            set_cell_margins(cell, top=55, bottom=55, left=80, right=80)

    set_table_borders_bw(tbl_scen)

    # ---------------------------------------------------------------------------
    # BAB VI: SPRINT REVIEW & RETROSPECTIVE
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "VI. SPRINT REVIEW DAN SPRINT RETROSPECTIVE", level=1)

    add_styled_heading(doc, "6.1 Checklist Sprint Review P7", level=2)
    tbl_rev = doc.add_table(rows=8, cols=2)
    tbl_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rev.autofit = False

    rev_headers = ["Item Evaluasi Sprint Review", "Status Hasil"]
    for i, h in enumerate(rev_headers):
        cell = tbl_rev.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    rev_rows = [
        ("Abstract class Produk berhasil dibuat", "TERVERIFIKASI (100%)"),
        ("Abstract method getKategoriInfo() & tampilkanDetailKhusus() diimplementasikan subclass", "TERVERIFIKASI (100%)"),
        ("Minimal 2 subclass konkret (ProdukMakanan & ProdukElektronik) dapat digunakan", "TERVERIFIKASI (100%)"),
        ("Interface DapatDidiskon berhasil diimplementasikan oleh kedua subclass", "TERVERIFIKASI (100%)"),
        ("Class diagram diperbarui sesuai notasi UML dan konsisten dengan kode program", "TERVERIFIKASI (100%)"),
        ("Pengujian polymorphism melalui superclass dan interface berhasil 100%", "TERVERIFIKASI (100%)"),
        ("Fitur proyek sebelumnya (P1-P6: Association, Composition, Overloading) tetap berjalan", "TERVERIFIKASI (100%)")
    ]

    r_widths = [Inches(4.5), Inches(2.0)]
    for r_idx, (c0, c1) in enumerate(rev_rows, start=1):
        row = tbl_rev.rows[r_idx]
        for c_idx, val in enumerate([c0, c1]):
            cell = row.cells[c_idx]
            cell.width = r_widths[c_idx]
            p = cell.paragraphs[0]
            if c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if c_idx == 1:
                run.font.bold = True
            set_cell_margins(cell, top=55, bottom=55, left=100, right=100)

    set_table_borders_bw(tbl_rev)

    add_styled_heading(doc, "6.2 Sprint Retrospective P7", level=2)
    add_bullet_p(doc, "Transformasi superclass Produk menjadi abstract class berjalan mulus tanpa merusak relasi Aggregation pada ItemTransaksi. "
                      "Realisasi interface DapatDidiskon pada subclass ProdukMakanan dan ProdukElektronik berhasil menyediakan standarisasi kontrak diskon promosi "
                      "yang sangat fleksibel dan terbukti polimorfik.", bold_prefix="What Went Well? ")
    add_bullet_p(doc, "Diperlukan pertimbangan mendalam dalam membedakan method yang sebaiknya dijadikan abstract method murni versus method konkret "
                      "dengan Template Method Pattern. Solusi yang diambil adalah mempertahankan logika umum cetak data di Produk dan mendelegasikan detail khusus ke subclass.", bold_prefix="What Went Wrong? ")
    add_bullet_p(doc, "Proyek hasil P7 ini telah siap menjadi baseline solid untuk menghadapi P8 (UTS Praktikum OOP) dan P9 (Penyempurnaan UML Class Diagram "
                      "ke struktur program Java enterprise lanjutan).", bold_prefix="Improvement: ")

    # ---------------------------------------------------------------------------
    # BAB VII: DEFINITION OF DONE (15/15) & KESIMPULAN
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "VII. DEFINITION OF DONE DAN KESIMPULAN", level=1)
    
    add_body_p(doc, "Berdasarkan pedoman evaluasi Modul 7, seluruh 15 kriteria Definition of Done telah terpenuhi secara paripurna:")
    dods = [
        "Menggunakan proyek hasil P6 (tidak membuat proyek dari nol).",
        "Audit desain dan alasan penggunaan abstraksi tersedia secara tertulis.",
        "Minimal 1 abstract class yang relevan dibuat (Produk.java).",
        "Minimal 1 abstract method dibuat (getKategoriInfo() dan tampilkanDetailKhusus()).",
        "Minimal 2 subclass konkret mengimplementasikan abstract method (ProdukMakanan & ProdukElektronik).",
        "Minimal 1 interface dibuat (DapatDidiskon.java).",
        "Minimal 2 class mengimplementasikan interface yang relevan.",
        "Polymorphism melalui superclass dan interface dapat dibuktikan melalui runtime dispatch.",
        "Class diagram diperbarui dan konsisten dengan implementasi kode sumber.",
        "Program berhasil dikompilasi (javac) dan dijalankan (java) tanpa error.",
        "Minimal 4 skenario pengujian tersedia (terpenuhi 5 skenario komprehensif).",
        "Fitur proyek sebelumnya tetap berjalan utuh tanpa regresi.",
        "Sprint Backlog P7 diperbarui dan berstatus DONE.",
        "Sprint Review P7 tersedia dan terverifikasi 100%.",
        "Sprint Retrospective P7 tersedia dengan analisis mendalam."
    ]
    for d in dods:
        add_bullet_p(doc, d, bold_prefix="[V] ")

    add_body_p(doc, 
               "Dengan selesainya seluruh tahapan pada Sprint P7 ini, arsitektur Sistem Kasir Sederhana telah memiliki pembagian tanggung jawab "
               "yang sangat jelas, hierarki class yang kokoh, serta kontrak perilaku yang konsisten dan teruji polimorfik.")

    doc.save(output_filename)
    print(f"Laporan Word saved at: {output_filename}")
    return output_filename

if __name__ == "__main__":
    docx_path1 = os.path.join(base_dir, "Laporan_P7_3125522007_Rafael_Rizky.docx")
    docx_path2 = os.path.join(base_dir, "README.docx")
    build_laporan_p7(docx_path1)
    build_laporan_p7(docx_path2)
    print("All Word reports created successfully.")

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_paragraph(p, font_name="Times New Roman", font_size=12, bold=False, italic=False, 
                     color_rgb=(0,0,0), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line_spacing=1.15):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 30, 30)
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(50, 50, 50)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11.5)
    r_body.font.italic = italic
    r_body.font.color.rgb = RGBColor(20, 20, 20)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(20, 20, 20)
    return p

def build_laporan_docx(output_filename):
    doc = docx.Document()

    # Set page margins: 1 inch (72 pt / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set base style to Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # HEADER / TITLE
    # ---------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("LAPORAN PRAKTIKUM PEMROGRAMAN BERORIENTASI OBYEK")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r = p_sub.add_run("MODUL 2: IMPLEMENTASI CLASS, OBJECT, ATTRIBUTE, METHOD, DAN CONSTRUCTOR")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    # Identity box / table
    table_id = doc.add_table(rows=4, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False
    
    id_data = [
        ("Nama Mahasiswa", ": Rafael Rizky"),
        ("NRP", ": 3125522007"),
        ("Dosen Pengampu", ": Nirwana Haidar Hari, S.Pd., M.Kom."),
        ("Program Studi / Kampus", ": D3 Teknik Informatika - PENS PSDKU Sumenep (2026)")
    ]
    for idx, (label, val) in enumerate(id_data):
        row = table_id.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        r0.font.bold = True
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------------------------
    # HALAMAN 1 / BAGIAN 1: REVIEW PROYEK P1 & AGILE SPRINT P2
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "1. Deskripsi Proyek dan Review Hasil P1", level=1)
    
    add_body_p(doc, "Sistem Kasir Sederhana", bold_prefix="1.1 Nama Proyek: ")
    add_body_p(doc, "Proses pencatatan produk dan transaksi penjualan pada toko ritel/kuliner memerlukan pengelolaan yang terstruktur dan teratur. Pencatatan konvensional secara manual berisiko tinggi menimbulkan ketidaksesuaian data harga, kekeliruan perhitungan total transaksi, serta sulitnya memantau sisa stok barang secara akurat. Sistem Kasir Sederhana ini dikembangkan untuk mengotomatisasi pencatatan data produk, mengelola data pelanggan, serta memfasilitasi transaksi penjualan secara terintegrasi.", bold_prefix="1.2 Latar Belakang Masalah: ")
    add_body_p(doc, "Membuat aplikasi kasir sederhana berbasis console untuk membantu mengelola data produk, memelihara informasi pelanggan, serta mendukung proses kalkulasi dan pencatatan transaksi penjualan secara akurat dan teratur.", bold_prefix="1.3 Product Goal: ")

    add_body_p(doc, "Aktor dan Peran dalam Sistem", bold_prefix="1.4 ")
    t_aktor = doc.add_table(rows=3, cols=2)
    t_aktor.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_aktor)
    t_aktor_data = [
        ("Aktor", "Peran dalam Sistem"),
        ("Kasir", "Mengelola dan menggunakan data produk dalam transaksi penjualan, melayani pembelian pelanggan, serta mencetak struk transaksi."),
        ("Pemilik / Admin", "Mengelola master data produk (menambah stok, memperbarui harga, memantau nilai inventaris), serta mengelola data pelanggan/membership.")
    ]
    for r_idx, row_content in enumerate(t_aktor_data):
        for c_idx, val in enumerate(row_content):
            cell = t_aktor.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.8)
            else:
                cell.width = Inches(4.7)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            if r_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "EAEAEA")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_body_p(doc, "Product Backlog Awal (Hasil P1)", bold_prefix="1.5 ")
    t_pb = doc.add_table(rows=6, cols=3)
    t_pb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_pb)
    pb_data = [
        ("ID", "User Story", "Prioritas"),
        ("US-01", "Sebagai admin, saya ingin menambahkan data produk, sehingga data produk dapat dikelola.", "High"),
        ("US-02", "Sebagai kasir, saya ingin melihat daftar produk, sehingga saya dapat mengetahui produk yang tersedia.", "High"),
        ("US-03", "Sebagai admin, saya ingin mengubah data produk, sehingga informasi produk tetap sesuai.", "Medium"),
        ("US-04", "Sebagai kasir, saya ingin memproses transaksi penjualan, sehingga transaksi dapat dicatat.", "High"),
        ("US-05", "Sebagai admin, saya ingin melihat data stok produk, sehingga dapat mengetahui ketersediaan produk.", "Medium")
    ]
    for r_idx, row_content in enumerate(pb_data):
        for c_idx, val in enumerate(row_content):
            cell = t_pb.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(0.9)
            elif c_idx == 1:
                cell.width = Inches(4.7)
            else:
                cell.width = Inches(0.9)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            if r_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "EAEAEA")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Sprint Planning P2
    add_styled_heading(doc, "2. Perencanaan Agile Sprint (Sprint P2)", level=1)
    add_body_p(doc, "Mengimplementasikan class utama proyek (Produk, Pelanggan, dan Transaksi) ke dalam kode pemrograman Java secara utuh, sehingga setiap object dapat diinstansiasi, diinisialisasi menggunakan constructor, menjalankan manipulasi atribut melalui method berparameter, serta menghasilkan kalkulasi nilai kembalian (return value).", bold_prefix="2.1 Sprint Goal: ")
    
    add_body_p(doc, "Sprint Backlog P2", bold_prefix="2.2 ")
    t_sb = doc.add_table(rows=7, cols=3)
    t_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sb)
    sb_data = [
        ("ID", "Sprint Backlog", "Status"),
        ("SB-01", "Mendefinisikan dan membuat class Produk beserta atribut dan constructor.", "DONE"),
        ("SB-02", "Mendefinisikan dan membuat class Pelanggan beserta atribut dan constructor.", "DONE"),
        ("SB-03", "Mendefinisikan dan membuat class Transaksi beserta atribut relasi dan constructor.", "DONE"),
        ("SB-04", "Mengimplementasikan method tanpa parameter pada seluruh class (tampilkanData, prosesTransaksi).", "DONE"),
        ("SB-05", "Mengimplementasikan method dengan parameter dan method dengan return value pada seluruh class.", "DONE"),
        ("SB-06", "Mengimplementasikan class Main untuk instansiasi minimal 2 object per class serta pengujian fitur.", "DONE")
    ]
    for r_idx, row_content in enumerate(sb_data):
        for c_idx, val in enumerate(row_content):
            cell = t_sb.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.0)
            elif c_idx == 1:
                cell.width = Inches(4.5)
            else:
                cell.width = Inches(1.0)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            if r_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "EAEAEA")
            elif c_idx == 2:
                run.font.bold = True

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 2 / BAGIAN 2: REFINEMENT OBJECT & CLASS DIAGRAM
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "3. Refinement Object dan Perancangan Class", level=1)
    add_body_p(doc, "Pada tahapan ini dilakukan pendefinisian secara rinci terhadap struktur masing-masing class yang diturunkan dari kandidat object pada P1. Setiap class dirancang agar memiliki minimal 3 atribut, 1 constructor untuk inisialisasi, serta kombinasi method operasional yang mencakup method tanpa parameter, method dengan parameter, dan method yang mengembalikan nilai.")

    add_body_p(doc, "Tabel Refinement Object dan Perancangan Class", bold_prefix="3.1 ")
    t_ref = doc.add_table(rows=4, cols=4)
    t_ref.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_ref)
    ref_data = [
        ("Nama Class", "Attribute", "Method", "Constructor"),
        ("Produk", 
         "• kode : String\n• nama : String\n• harga : int\n• stok : int",
         "• tampilkanData() : void\n• ubahHarga(int) : void\n• tambahStok(int) : void\n• kurangiStok(int) : void\n• hitungNilaiInventaris() : int\n• getNama() : String",
         "Produk(String kode, String nama, int harga, int stok)"),
        ("Pelanggan",
         "• idPelanggan : String\n• nama : String\n• nomorHP : String\n• tipeMember : String",
         "• tampilkanData() : void\n• ubahNomorHP(String) : void\n• ubahTipeMember(String) : void\n• getDiskon() : double\n• getNama() : String",
         "Pelanggan(String idPelanggan, String nama, String nomorHP, String tipeMember)"),
        ("Transaksi",
         "• idTransaksi : String\n• tanggal : String\n• pelanggan : Pelanggan\n• produk : Produk\n• jumlahBeli : int",
         "• hitungSubtotal() : int\n• hitungDiskonNominal() : double\n• hitungTotalBayar() : double\n• ubahJumlahBeli(int) : void\n• prosesTransaksi() : void",
         "Transaksi(String idTransaksi, String tanggal, Pelanggan pelanggan, Produk produk, int jumlahBeli)")
    ]
    for r_idx, row_content in enumerate(ref_data):
        for c_idx, val in enumerate(row_content):
            cell = t_ref.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.1)
            elif c_idx == 1:
                cell.width = Inches(1.6)
            elif c_idx == 2:
                cell.width = Inches(2.2)
            else:
                cell.width = Inches(1.6)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            if r_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "EAEAEA")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_body_p(doc, "Diagram Sederhana Class (UML Class Diagram)", bold_prefix="3.2 ")
    add_body_p(doc, "Hubungan antar class digambarkan melalui Class Diagram berikut. Class Transaksi mengasosiasikan class Pelanggan dan class Produk untuk merekam entitas pembeli dan produk yang dibeli beserta penghitungan potongan diskon member.")

    diagram_path = os.path.join(r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\assets", "class_diagram.png")
    if os.path.exists(diagram_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(diagram_path, width=Inches(6.2))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("Gambar 1. UML Class Diagram Sistem Kasir Sederhana")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True

    add_body_p(doc, "Penjelasan Hubungan dan Desain Method:", bold_prefix="3.3 ")
    add_bullet_p(doc, "Menyimpan data item berupa kode, nama, harga satuan, dan jumlah stok. Memiliki method ubahHarga() dan tambahStok() dengan parameter untuk modifikasi data, serta hitungNilaiInventaris() dengan nilai kembalian (harga * stok).", bold_prefix="1. Class Produk: ")
    add_bullet_p(doc, "Menyimpan informasi pembeli, kontak, dan klasifikasi membership (VIP, Gold, Reguler). Memiliki method getDiskon() dengan return value double yang memberikan potongan 15% untuk VIP dan 10% untuk Gold.", bold_prefix="2. Class Pelanggan: ")
    add_bullet_p(doc, "Mengintegrasikan objek Pelanggan dan Produk. Memiliki kalkulasi hitungSubtotal() dan hitungTotalBayar() (return value), fasilitas ubahJumlahBeli(int), serta prosesTransaksi() yang mencetak nota kasir sekaligus memotong stok produk secara otomatis.", bold_prefix="3. Class Transaksi: ")

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 3 / BAGIAN 3: PENGUJIAN, HASIL RUNNING, REVIEW & RETROSPECTIVE
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "4. Pengujian Program dan Hasil Eksekusi", level=1)
    add_body_p(doc, "Pengujian dilakukan melalui class Main.java dengan skenario pengujian komprehensif yang mencakup pembuktian seluruh kriteria Modul 2:")
    add_bullet_p(doc, "Instansiasi minimal 2 object untuk setiap class (produk1 & produk2, pelanggan1 & pelanggan2, transaksi1 & transaksi2) dengan inisialisasi constructor.", bold_prefix="a. Instansiasi Object: ")
    add_bullet_p(doc, "Memanggil method getNama(), hitungNilaiInventaris(), dan getDiskon() untuk membaca kalkulasi nilai.", bold_prefix="b. Pengujian Return Value: ")
    add_bullet_p(doc, "Mengubah harga produk1 (ubahHarga(16000)), menambah stok produk1 (tambahStok(10)), memperbarui kontak (ubahNomorHP), serta menaikkan level pelanggan2 menjadi Gold.", bold_prefix="c. Perubahan Data (Method Berparameter): ")
    add_bullet_p(doc, "Memanggil method prosesTransaksi() untuk TRX-001 dan TRX-002, memeriksa pengurangan stok produk, dan mencetak struk transaksi.", bold_prefix="d. Proses Transaksi & Struk Kasir: ")

    add_body_p(doc, "Tangkapan Layar Hasil Running Program", bold_prefix="4.1 ")
    term_img_path = os.path.join(r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\assets", "hasil_running.png")
    if os.path.exists(term_img_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(4)
        p_img2.paragraph_format.space_after = Pt(2)
        run_img2 = p_img2.add_run()
        run_img2.add_picture(term_img_path, width=Inches(6.0))

        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_before = Pt(0)
        p_cap2.paragraph_format.space_after = Pt(8)
        r_cap2 = p_cap2.add_run("Gambar 2. Bukti Running Program Main.java pada Terminal Windows PowerShell")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10)
        r_cap2.font.italic = True

    add_styled_heading(doc, "5. Sprint Review (Evaluasi Sprint P2)", level=1)
    add_body_p(doc, "Setelah seluruh implementasi dan pengujian diselesaikan, evaluasi ketercapaian sprint dirangkum pada tabel berikut:")
    
    t_rev = doc.add_table(rows=7, cols=2)
    t_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_rev)
    rev_data = [
        ("Item Evaluasi", "Hasil / Keterangan"),
        ("Class berhasil dibuat", "3 class utama (Produk, Pelanggan, Transaksi) dan 1 class runner (Main) berhasil dibuat sesuai spesifikasi."),
        ("Object berhasil dibuat", "Berhasil diinstansiasi minimal 2 object untuk setiap class (total 6 object aktif)."),
        ("Constructor berjalan", "Constructor berparameter pada setiap class sukses menginisialisasi atribut object."),
        ("Method berjalan", "Seluruh method tanpa parameter, dengan parameter, dan method ber-return value berfungsi normal tanpa kendala."),
        ("Program dapat dijalankan", "Program berhasil dikompilasi menggunakan javac 26.0.1 dan dijalankan tanpa runtime exception."),
        ("Kendala", "Tidak ditemukan kendala fatal. Penyesuaian sintaks wildcard kompilasi pada shell Windows PowerShell diselesaikan dengan pemetaan berkas secara dinamis.")
    ]
    for r_idx, row_content in enumerate(rev_data):
        for c_idx, val in enumerate(row_content):
            cell = t_rev.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(2.2)
            else:
                cell.width = Inches(4.3)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            if r_idx == 0:
                run.font.bold = True
                set_cell_background(cell, "EAEAEA")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "6. Sprint Retrospective", level=1)
    add_body_p(doc, "Seluruh target sprint berhasil dicapai dengan tepat waktu. Struktur kode class Produk, Pelanggan, dan Transaksi dirancang modular, clean, serta saling berinteraksi secara harmonis dalam mendukung skenario transaksi kasir.", bold_prefix="What Went Well? ")
    add_body_p(doc, "Atribut pada setiap class saat ini masih berstatus package-private/default sehingga rentan dimodifikasi secara langsung dari luar class tanpa melalui validasi logika bisnis yang ketat.", bold_prefix="What Went Wrong? ")
    add_body_p(doc, "Pada sprint berikutnya (P3), struktur class akan ditingkatkan dengan menerapkan prinsip Encapsulation, yaitu mengubah akses atribut menjadi private serta menyediakan method Getter, Setter, dan validasi data terstruktur.", bold_prefix="Improvement: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "7. Definition of Done (DoD) Checklist", level=1)
    dod_items = [
        ("Menggunakan proyek P1 (Sistem Kasir Sederhana)", True),
        ("Minimal 3 class dibuat (Produk, Pelanggan, Transaksi)", True),
        ("Setiap class memiliki minimal 3 attribute (masing-masing 4–5 atribut)", True),
        ("Setiap class memiliki constructor untuk inisialisasi data awal", True),
        ("Setiap class memiliki minimal 2 method utama", True),
        ("Terdapat method dengan parameter (misal: ubahHarga, tambahStok, ubahJumlahBeli)", True),
        ("Terdapat method dengan return value (misal: hitungNilaiInventaris, getDiskon, hitungTotalBayar)", True),
        ("Minimal 2 object dibuat dari setiap class pada Main.java", True),
        ("Seluruh object dapat digunakan dan berinteraksi secara terintegrasi dari Main", True),
        ("Program berhasil dikompilasi bebas dari error", True),
        ("Program berhasil dijalankan dan output sesuai dengan ekspektasi", True),
        ("Sprint Backlog diperbarui dengan status DONE", True),
        ("Sprint Review dan Sprint Retrospective telah disusun secara lengkap", True)
    ]
    t_dod = doc.add_table(rows=len(dod_items)+1, cols=2)
    t_dod.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_dod)
    
    hdr0, hdr1 = t_dod.rows[0].cells[0], t_dod.rows[0].cells[1]
    hdr0.width, hdr1.width = Inches(5.3), Inches(1.2)
    set_cell_margins(hdr0, 80, 80, 100, 100)
    set_cell_margins(hdr1, 80, 80, 100, 100)
    set_cell_background(hdr0, "EAEAEA")
    set_cell_background(hdr1, "EAEAEA")
    p = hdr0.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Kriteria Definition of Done (Modul 2)"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(10.5)
    p = hdr1.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Status"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(10.5)

    for idx, (item, fulfilled) in enumerate(dod_items, start=1):
        c0, c1 = t_dod.rows[idx].cells[0], t_dod.rows[idx].cells[1]
        c0.width, c1.width = Inches(5.3), Inches(1.2)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0); p0.paragraph_format.line_spacing = 1.1
        r0 = p0.add_run(item); r0.font.name = "Times New Roman"; r0.font.size = Pt(10)
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0); p1.paragraph_format.line_spacing = 1.1
        r1 = p1.add_run("[ v ] Terpenuhi" if fulfilled else "[   ] Belum"); r1.font.name = "Times New Roman"; r1.font.size = Pt(10); r1.font.bold = True

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # LAMPIRAN: SOURCE CODE JAVA LENGKAP
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "Lampiran: Source Code Java", level=1)
    
    code_files = [
        ("Produk.java", r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\src\Produk.java"),
        ("Pelanggan.java", r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\src\Pelanggan.java"),
        ("Transaksi.java", r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\src\Transaksi.java"),
        ("Main.java", r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\src\Main.java")
    ]

    for fname, fpath in code_files:
        add_styled_heading(doc, f"Berkas: {fname}", level=2)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                code_content = f.read()
            
            # Put in a single-cell table for clean presentation
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(6.5)
            set_cell_background(cell, "F9F9F9")
            set_cell_margins(cell, 120, 120, 150, 150)
            set_table_borders(tbl)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(code_content)
            r.font.name = "Courier New"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 30, 30)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_filename)
    print(f"Document saved successfully: {output_filename}")

out_docx = r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\Laporan_P2_3125522007_Rafael_Rizky.docx"
build_laporan_docx(out_docx)

# Also save as README.docx in the same folder for standard project reference
readme_docx = r"d:\KULIAH\Semester 3\OOP\P2_3125522007_Rafael_Rizky\README.docx"
build_laporan_docx(readme_docx)

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_bw(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11.5)
    r_body.font.italic = italic
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_laporan_p4(output_filename):
    doc = docx.Document()

    # Margin 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # ---------------------------------------------------------------------------
    # HEADER / TITLE (Pure Black and White)
    # ---------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("LAPORAN PRAKTIKUM PEMROGRAMAN BERORIENTASI OBYEK")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r = p_sub.add_run("MODUL 4: RELASI ANTAROBJECT (ASSOCIATION, AGGREGATION, DAN COMPOSITION)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    # Identity Table (Black and White)
    table_id = doc.add_table(rows=4, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False
    
    id_data = [
        ("Nama Mahasiswa", ": Rafael Rizky"),
        ("NRP", ": 3125522007"),
        ("Dosen Pengampu", ": Nirwana Haidar Hari, S.Pd., M.Kom."),
        ("Program Studi / Kampus", ": D3 Teknik Informatika - PENS PSDKU Sumenep (2026)")
    ]
    for idx, (label, val) in enumerate(id_data):
        row = table_id.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0, 0, 0)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------------------------
    # HALAMAN 1 / BAGIAN 1: REVIEW P3, SPRINT PLANNING & IDENTIFIKASI RELASI
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "1. Deskripsi Proyek dan Review Hasil P3", level=1)
    add_body_p(doc, "Sistem Kasir Sederhana", bold_prefix="1.1 Nama Proyek: ")
    add_body_p(doc, "Pada praktikum P3, class Produk, Pelanggan, dan Transaksi telah berhasil dienkapsulasi dengan pengamanan private modifier, getter, setter, serta validasi data. Namun, class-class tersebut masih berdiri sendiri dan relasi antarobjek belum dimodelkan secara komprehensif. Pada transaksi kasir di dunia nyata, sebuah transaksi dapat memuat banyak variasi barang belanjaan (multi-item), di mana data produk bersifat mandiri, data pelanggan berinteraksi sebagai pembeli, dan rincian belanja merupakan bagian tak terpisahkan dari transaksi. Oleh karena itu, pada praktikum P4 ini dilakukan pengembangan relasi antarobject mencakup Association, Aggregation, dan Composition.", bold_prefix="1.2 Latar Belakang Masalah: ")

    add_styled_heading(doc, "2. Perencanaan Agile Sprint (Sprint P4)", level=1)
    add_body_p(doc, "Menghubungkan class utama proyek sehingga seluruh object dapat saling berinteraksi sebagai satu kesatuan sistem kasir yang utuh melalui penerapan relasi Association, Aggregation, dan Composition.", bold_prefix="2.1 Sprint Goal: ")

    add_body_p(doc, "Sprint Backlog P4", bold_prefix="2.2 ")
    t_sb = doc.add_table(rows=7, cols=3)
    t_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_sb)
    sb_data = [
        ("ID", "Sprint Backlog", "Status"),
        ("SB-01", "Mengidentifikasi relasi antarclass (Association, Aggregation, Composition).", "DONE"),
        ("SB-02", "Memperbarui UML Class Diagram proyek dengan notasi relasi standar.", "DONE"),
        ("SB-03", "Mengimplementasikan relasi Association antara Pelanggan dan Transaksi.", "DONE"),
        ("SB-04", "Mengimplementasikan relasi Aggregation antara ItemTransaksi dan Produk.", "DONE"),
        ("SB-05", "Mengimplementasikan relasi Composition antara Transaksi dan ItemTransaksi.", "DONE"),
        ("SB-06", "Membuat program pengujian interaksi multi-objek pada Main.java.", "DONE")
    ]
    for r_idx, row_content in enumerate(sb_data):
        for c_idx, val in enumerate(row_content):
            cell = t_sb.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(0.9)
            elif c_idx == 1:
                cell.width = Inches(4.7)
            else:
                cell.width = Inches(0.9)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0 or c_idx == 2:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_styled_heading(doc, "3. Identifikasi Relasi Antarclass (Bagian A)", level=1)
    add_body_p(doc, "Sistem Kasir Sederhana menggunakan 4 class utama: Produk, Pelanggan, ItemTransaksi, dan Transaksi. Analisis relasi antarclass dijabarkan pada tabel di bawah ini:")

    t_rel = doc.add_table(rows=4, cols=4)
    t_rel.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_rel)
    rel_data = [
        ("Class A", "Class B", "Relasi", "Alasan & Derajat Ketergantungan"),
        ("Pelanggan", "Transaksi", "Association", "Hubungan 'uses-a' (lemah). Pelanggan berinteraksi melakukan transaksi belanja. Objek Pelanggan tetap hidup mandiri di luar transaksi."),
        ("ItemTransaksi", "Produk", "Aggregation", "Hubungan 'has-a' (sedang). ItemTransaksi mereferensikan Produk. Objek Produk tetap eksis secara mandiri pada katalog/stok toko meskipun transaksi selesai."),
        ("Transaksi", "ItemTransaksi", "Composition", "Hubungan 'part-of' (kuat). ItemTransaksi diciptakan dan dikelola di dalam Transaksi. Jika Transaksi dihapus, ItemTransaksi ikut musnah.")
    ]
    for r_idx, row_content in enumerate(rel_data):
        for c_idx, val in enumerate(row_content):
            cell = t_rel.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.2)
            elif c_idx == 1:
                cell.width = Inches(1.2)
            elif c_idx == 2:
                cell.width = Inches(1.2)
            else:
                cell.width = Inches(2.9)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 2 / BAGIAN 2: UML CLASS DIAGRAM & PENJELASAN RELASI
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "4. Update Class Diagram (Bagian B)", level=1)
    add_body_p(doc, "Struktur relasi antarclass digambarkan melalui diagram kelas UML monokrom (hitam-putih) yang memperlihatkan notasi association, aggregation (berlian kosong), dan composition (berlian penuh solid) beserta tingkat multiplicity:")

    diagram_path = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\assets", "class_diagram.png")
    if os.path.exists(diagram_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(diagram_path, width=Inches(6.2))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("Gambar 1. UML Class Diagram Relasi Antarobject Sistem Kasir Sederhana (P4)")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0, 0, 0)

    add_styled_heading(doc, "5. Penjelasan Implementasi Relasi Antarobject (Bagian C, D, E)", level=1)
    
    add_body_p(doc, "Association merepresentasikan relasi struktural di mana objek saling berinteraksi tanpa adanya kepemilikan siklus hidup. Pada sistem ini, class Transaksi memiliki atribut bertipe Pelanggan ('private Pelanggan pelanggan;'). Objek Pelanggan dibuat secara independen di luar dan dihubungkan ke Transaksi untuk menghitung besaran diskon member (VIP 15%, GOLD 10%, REGULER 0%) serta mencetak nama pembeli pada struk.", bold_prefix="5.1 Implementasi Association (Pelanggan - Transaksi): ")
    
    add_body_p(doc, "Aggregation merupakan hubungan 'has-a' di mana objek bagian dapat hidup mandiri tanpa objek induknya. Pada sistem kasir, class ItemTransaksi mengagregasikan class Produk ('private Produk produk;'). Objek Produk dibuat di luar (pada master katalog produk) dan dimasukkan sebagai referensi ke dalam ItemTransaksi. Apabila transaksi selesai atau dibatalkan, objek Produk tidak ikut musnah, melainkan tetap eksis di dalam inventaris toko.", bold_prefix="5.2 Implementasi Aggregation (ItemTransaksi - Produk): ")
    
    add_body_p(doc, "Composition merupakan hubungan kepemilikan kuat 'part-of' di mana siklus hidup objek bagian bergantung penuh pada objek pemiliknya. Pada sistem ini, class Transaksi mengomposisi kumpulan objek ItemTransaksi ('private ArrayList<ItemTransaksi> daftarItem;'). Objek ItemTransaksi diciptakan langsung di dalam method tambahItem() pada class Transaksi ('ItemTransaksi itemBaru = new ItemTransaksi(produk, jumlahBeli);'). Jika objek Transaksi dimusnahkan, maka seluruh objek ItemTransaksi di dalamnya otomatis ikut musnah.", bold_prefix="5.3 Implementasi Composition (Transaksi - ItemTransaksi): ")

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # HALAMAN 3 / BAGIAN 3: PENGUJIAN, SPRINT REVIEW & RETROSPECTIVE
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "6. Pengujian Interaksi Object (Bagian F) dan Hasil Running", level=1)
    add_body_p(doc, "Pengujian pada Main.java dirancang mencakup 3 skenario wajib sesuai panduan Modul 4:")
    add_bullet_p(doc, "Instansiasi mandiri objek Produk (P001, P002, P003), Pelanggan (VIP, GOLD, REGULER), dan Transaksi (TRX-001, TRX-002). Seluruh objek berhasil diinisialisasi tanpa error.", bold_prefix="a. Test 1 (Object Berhasil Dibuat): ")
    add_bullet_p(doc, "Objek Transaksi berinteraksi dengan Pelanggan (Association) dan dengan Produk melalui ItemTransaksi (Aggregation & Composition). Transaksi memvalidasi sisa stok saat item ditambahkan dan menolak pembelian yang melebihi kuota stok.", bold_prefix="b. Test 2 (Dua atau Lebih Object Berinteraksi): ")
    add_bullet_p(doc, "Objek Transaksi mengekstrak harga dari Produk via ItemTransaksi untuk kalkulasi subtotal, memanggil getDiskon() dari Pelanggan untuk potongan harga, serta mengurangi stok produk secara otomatis saat transaksi diproses.", bold_prefix="c. Test 3 (Pemanfaatan Data Object Lain Lewat Method): ")

    add_body_p(doc, "Tangkapan Layar Hasil Running Terminal (Monokrom)", bold_prefix="6.1 ")
    term_img_path = os.path.join(r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\assets", "hasil_running.png")
    if os.path.exists(term_img_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(4)
        p_img2.paragraph_format.space_after = Pt(2)
        run_img2 = p_img2.add_run()
        run_img2.add_picture(term_img_path, width=Inches(6.0))

        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_before = Pt(0)
        p_cap2.paragraph_format.space_after = Pt(8)
        r_cap2 = p_cap2.add_run("Gambar 2. Bukti Running Program Main.java pada Terminal Windows PowerShell")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(0, 0, 0)

    add_styled_heading(doc, "7. Sprint Review (Evaluasi Sprint P4)", level=1)
    t_rev = doc.add_table(rows=8, cols=2)
    t_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_rev)
    rev_data = [
        ("Item Evaluasi", "Hasil / Keterangan"),
        ("Class diagram diperbarui", "Class diagram berhasil diperbarui dengan 4 class dan notasi relasi lengkap."),
        ("Association berhasil", "Relasi association (Pelanggan - Transaksi) berhasil diterapkan dan diuji."),
        ("Aggregation berhasil", "Relasi aggregation (ItemTransaksi - Produk) sukses menghubungkan master produk."),
        ("Composition berhasil", "Relasi composition (Transaksi - ItemTransaksi) sukses mengelola multi-item."),
        ("Object dapat berinteraksi", "Seluruh objek saling bertukar pesan dan data secara teratur dan aman."),
        ("Program berjalan", "Program sukses dikompilasi (javac) dan dieksekusi (java) tanpa runtime error."),
        ("Kendala", "Tidak ditemukan kendala fatal. Penanganan multi-item diselesaikan menggunakan ArrayList.")
    ]
    for r_idx, row_content in enumerate(rev_data):
        for c_idx, val in enumerate(row_content):
            cell = t_rev.rows[r_idx].cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(2.2)
            else:
                cell.width = Inches(4.3)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "8. Sprint Retrospective", level=1)
    add_body_p(doc, "Penerapan ketiga jenis relasi (Association, Aggregation, dan Composition) berhasil diwujudkan secara nyata pada studi kasus kasir. Fitur multi-item belanja berjalan sangat baik dan struk kasir tercetak terstruktur.", bold_prefix="What Went Well? ")
    add_body_p(doc, "Penentuan batasan kepemilikan siklus hidup antara Aggregation dan Composition memerlukan analisis mendalam agar objek bagian tidak salah dideklarasikan.", bold_prefix="What Went Wrong? ")
    add_body_p(doc, "Pada sprint berikutnya (P5: Inheritance dan Generalization), class Produk dan Pelanggan dapat digeneralisasi untuk membentuk hierarki kelas induk dan turunan.", bold_prefix="Improvement: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "9. Definition of Done (DoD) Checklist Modul 4", level=1)
    dod_items = [
        ("Menggunakan project P3 (Sistem Kasir Sederhana)", True),
        ("Minimal 4 class digunakan (Produk, Pelanggan, ItemTransaksi, Transaksi)", True),
        ("Relasi antarclass telah dianalisis secara terperinci", True),
        ("Class diagram diperbarui dengan notasi relasi standar", True),
        ("Minimal 1 association diterapkan (Pelanggan - Transaksi)", True),
        ("Minimal 1 aggregation diterapkan (ItemTransaksi - Produk)", True),
        ("Minimal 1 composition diterapkan (Transaksi - ItemTransaksi)", True),
        ("Object berinteraksi melalui pemanggilan method", True),
        ("Tidak mengakses attribute private secara langsung", True),
        ("Program berhasil dikompilasi bebas error", True),
        ("Program berhasil dijalankan dengan output sesuai", True),
        ("Minimal 3 skenario pengujian tersedia (Test 1, Test 2, Test 3)", True),
        ("Sprint Backlog diperbarui dengan status DONE", True),
        ("Sprint Review dan Sprint Retrospective tersedia lengkap", True)
    ]
    t_dod = doc.add_table(rows=len(dod_items)+1, cols=2)
    t_dod.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(t_dod)
    
    hdr0, hdr1 = t_dod.rows[0].cells[0], t_dod.rows[0].cells[1]
    hdr0.width, hdr1.width = Inches(5.3), Inches(1.2)
    set_cell_margins(hdr0, 70, 70, 90, 90)
    set_cell_margins(hdr1, 70, 70, 90, 90)
    p = hdr0.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Kriteria Definition of Done (Modul 4)"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0, 0, 0)
    p = hdr1.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run("Status"); r.font.name = "Times New Roman"; r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0, 0, 0)

    for idx, (item, fulfilled) in enumerate(dod_items, start=1):
        c0, c1 = t_dod.rows[idx].cells[0], t_dod.rows[idx].cells[1]
        c0.width, c1.width = Inches(5.3), Inches(1.2)
        set_cell_margins(c0, 50, 50, 80, 80)
        set_cell_margins(c1, 50, 50, 80, 80)
        
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0); p0.paragraph_format.line_spacing = 1.1
        r0 = p0.add_run(item); r0.font.name = "Times New Roman"; r0.font.size = Pt(9.5); r0.font.color.rgb = RGBColor(0, 0, 0)
        
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(0); p1.paragraph_format.line_spacing = 1.1
        r1 = p1.add_run("[ v ] Terpenuhi" if fulfilled else "[   ] Belum"); r1.font.name = "Times New Roman"; r1.font.size = Pt(9.5); r1.font.bold = True; r1.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # LAMPIRAN: SOURCE CODE JAVA LENGKAP P4 (Black & White formatting)
    # ---------------------------------------------------------------------------
    add_styled_heading(doc, "Lampiran: Source Code Java (Relasi Antarobject)", level=1)
    
    code_files = [
        ("Produk.java", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\src\Produk.java"),
        ("Pelanggan.java", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\src\Pelanggan.java"),
        ("ItemTransaksi.java", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\src\ItemTransaksi.java"),
        ("Transaksi.java", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\src\Transaksi.java"),
        ("Main.java", r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\src\Main.java")
    ]

    for fname, fpath in code_files:
        add_styled_heading(doc, f"Berkas: {fname}", level=2)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                code_content = f.read()
            
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.width = Inches(6.5)
            set_cell_margins(cell, 100, 100, 120, 120)
            set_table_borders_bw(tbl)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(code_content)
            r.font.name = "Courier New"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0, 0, 0)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_filename)
    print(f"Document saved successfully: {output_filename}")

out_docx = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\Laporan_P4_3125522007_Rafael_Rizky.docx"
build_laporan_p4(out_docx)

readme_docx = r"d:\KULIAH\Semester 3\OOP\Prakatikum\P4_3125522007_Rafael_Rizky\README.docx"
build_laporan_p4(readme_docx)
